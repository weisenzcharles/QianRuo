from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT_DIR = Path("/Users/CharlesZhang/Documents/租赁系统")
DOCX_OUT = OUT_DIR / "租赁项目设计文档.docx"
MD_OUT = OUT_DIR / "租赁项目设计文档.md"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(91, 100, 112)
BLACK = RGBColor(0, 0, 0)
LIGHT_GRAY = "F2F4F7"
LIGHT_BLUE = "E8EEF5"
DOC_FONT = "PingFang SC"


def set_run_font(run, size=None, bold=None, color=None):
    run.font.name = DOC_FONT
    run._element.rPr.rFonts.set(qn("w:ascii"), DOC_FONT)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), DOC_FONT)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), DOC_FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    total = sum(widths)
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_grid = tbl.tblGrid
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        tbl_grid.append(grid_col)
    for row in table.rows:
        for idx, width in enumerate(widths):
            cell = row.cells[idx]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.tcW
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def style_doc(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = DOC_FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), DOC_FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), DOC_FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), DOC_FONT)
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = styles[name]
        style.font.name = DOC_FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), DOC_FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), DOC_FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), DOC_FONT)
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_para(doc, text="", style=None, bold=False, color=None, size=None, align=None, after=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(item)
        set_run_font(run)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(item)
        set_run_font(run)


def add_table(doc, headers, rows, widths=None, header_fill=LIGHT_GRAY):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    header = table.rows[0].cells
    for i, text in enumerate(headers):
        set_cell_shading(header[i], header_fill)
        p = header[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        set_run_font(r, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(text))
            set_run_font(r)
    if widths is None:
        widths = [int(9360 / len(headers))] * len(headers)
    set_table_widths(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_callout(doc, label, text):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F4F6F9")
    set_cell_margins(cell, top=110, bottom=110, start=160, end=160)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(f"{label}：")
    set_run_font(r, bold=True, color=DARK_BLUE)
    r2 = p.add_run(text)
    set_run_font(r2)
    set_table_widths(table, [9360])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def masthead(doc):
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    header = section.header.paragraphs[0]
    header.text = "租赁系统项目设计文档"
    set_run_font(header.runs[0], size=9, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.text = "内部设计资料"
    set_run_font(footer.runs[0], size=9, color=MUTED)

    add_para(doc, "项目设计文档", size=24, bold=True, color=BLACK, after=4)
    add_para(doc, "迁渃租赁闭环及库存协同系统", size=14, color=MUTED, after=14)
    rows = [
        ("文档版本", "V1.0 整理稿"),
        ("整理日期", date.today().isoformat()),
        ("资料来源", "租赁项目文件夹内 18 份 Word 文档"),
        ("文档用途", "用于产品评审、需求澄清、研发拆解、测试验收与后续迭代规划"),
    ]
    add_table(doc, ["项目", "说明"], rows, widths=[1900, 7460], header_fill=LIGHT_BLUE)
    add_callout(
        doc,
        "设计边界",
        "平台以订单管理、SN 库存、物流归还、仓库协同和会员工具为核心；库存周转市场、多仓联盟等功能坚持信息展示与站内协同定位，不直接参与交易、定价、担保或资金结算。",
    )


SOURCE_DOCS = [
    "1.0版-0-总后台.docx",
    "1.0版-1-栏目导航.docx",
    "1.0版-2-首页.docx",
    "ok-1.0版-3-订单管理.docx",
    "1.0版-4-收发趋势.docx",
    "ok-1.0版-5-仓库管理.docx",
    "ok-1.0版-6-商品管理-.docx",
    "1.0版-7-公众号.docx",
    "1.0版-8-个人中心.docx",
    "ok-1.0版-9-极速核销台.docx",
    "1.0版-10-仓库调货.docx",
    "1.0版-15-免费与收费与推广.docx",
    "1.1版-订单中转大厅.docx",
    "1.1版-库存周转市场---待定.docx",
    "1.2版-多仓联盟.docx",
    "1.2版-快递下单及涉及流程+微信小程序.docx",
    "1.3版-免押码.docx",
    "2.0版-支付宝租赁小程序.docx",
]


def build_docx():
    doc = Document()
    style_doc(doc)
    masthead(doc)

    doc.add_heading("1. 项目概述", level=1)
    add_para(
        doc,
        "本项目是面向租赁商家的租赁业务管理平台，核心目标是把订单录入、库存占用、SN 流转、快递下单、客户归还、仓库协同、售后完结和经营统计做成一套可追溯的闭环系统。系统同时支持电脑端、微信端能力，以及后续支付宝租赁小程序、多仓联盟和库存周转市场扩展。",
    )
    add_table(
        doc,
        ["维度", "设计结论"],
        [
            ("核心用户", "租赁商家、仓库人员、企业会员/个人会员、平台运营管理员、联盟成员、推广员、终端归还客户。"),
            ("核心对象", "订单、商品、套餐、SN 设备、配件、仓库、物流单、视频/图片凭证、会员、账务记录、联盟协助订单、周转市场帖子。"),
            ("核心闭环", "录单前校验库存与时效，发货时绑定 SN 与物流，租期内跟踪状态，归还时验机入库，售后赔付或完结后更新库存与账务。"),
            ("平台原则", "基础功能免费或低门槛，主要通过增值服务、快递路由、免押、存储和推广体系产生收入；避免参与同行交易定价。"),
        ],
        widths=[1800, 7560],
    )

    doc.add_heading("2. 原始文档范围", level=1)
    add_para(doc, "本设计文档基于以下 18 份原始材料整理，已按模块与版本进行归并：")
    add_table(
        doc,
        ["序号", "文档名称", "归并模块"],
        [(i, name, classify_source(name)) for i, name in enumerate(SOURCE_DOCS, 1)],
        widths=[700, 4360, 4300],
    )

    doc.add_heading("3. 产品定位与业务边界", level=1)
    add_bullets(
        doc,
        [
            "平台主线是租赁商家内部经营系统，而不是开放式租赁交易平台。",
            "订单、库存、仓库、物流与验机之间必须形成强关联，所有关键动作保留日志。",
            "库存周转市场仅做库存展示、同行站内沟通与意向登记，不展示租金，不做订单、不做支付、不做担保。",
            "多仓联盟只做成员协助订单与记账结算信息，不在线支付，不替代成员之间的线下合作约定。",
            "客户侧微信小程序聚焦归还、订单查询、发货视频查看和开箱视频上传，不要求客户注册为平台会员。",
        ],
    )
    add_callout(
        doc,
        "合规策略",
        "涉及同行协作与周转市场的页面应默认隐藏手机号、微信、详细地址、真实姓名和自由文本导流信息；涉及客户订单的后台查看应按角色脱敏。",
    )

    doc.add_heading("4. 用户角色与权限", level=1)
    add_table(
        doc,
        ["角色", "主要能力", "权限/限制"],
        [
            ("平台总管理员", "会员、财务、库存、订单、微信小程序、授权记录、转发统计与异常账单管理。", "可查看平台级汇总；客户敏感信息仅管理员可查看。"),
            ("会员主账号", "管理仓库、商品、套餐、订单、权限组、复制/转发模板、视频保存策略、快递余额和增值服务。", "可设置库存是否对外接单，可注销账号。"),
            ("仓库人员", "查看指定仓库订单、备货、上传视频、填写单号、核销、验机、盘点。", "按会员配置的群组权限限定可见仓库和操作仓库。"),
            ("联盟成员", "查看联盟共享 SN 库存，接收或转派联盟协助订单，设置可承接上限。", "订单信息脱敏展示，只做记账结算价，不在线支付。"),
            ("终端客户", "通过微信小程序验证手机号后查询订单、查看发货视频和路由、提交归还快递、上传开箱视频。", "不可改归还仓库；手机号验证仅用于匹配订单。"),
            ("推广员", "使用推广码邀请会员，享受有效用户奖励、佣金和优惠包资格。", "需满足邀请、在线时长、真实订单、主体唯一、VIP 等申请条件。"),
        ],
        widths=[1500, 4560, 3300],
    )

    doc.add_heading("5. 信息架构与导航", level=1)
    add_table(
        doc,
        ["一级栏目", "核心内容"],
        [
            ("首页", "今日待发、今日归还、今日到库、汇总报表、营收分析、租用中、租赁地图。"),
            ("订单管理", "添加订单、所有订单、今日/明日/后日发货与归还、逾期、已发货、完成订单、售后、取消/追回、调货建议。"),
            ("收发趋势", "每日发货/归还趋势，发货数与收货数用不同颜色展示，商品/套餐列固定。"),
            ("库存概况/仓库管理", "在库空闲、预计库存、物流库存、仓库分类、告警配置、库存明细穿透。"),
            ("商品管理", "商品、配件、套餐、SN 档案、二维码/条码、履历、收益统计、出售与合作信息。"),
            ("设置中心/个人中心", "会员资料、仓库、产品套餐、权限、视频保存、时效查询、在线表格、记账簿、二维码工具、记事本。"),
            ("知识库", "品牌知识、通用知识、公告弹窗与确认。"),
            ("复购管理", "商品复购率、复购客户列表与关联订单。"),
            ("联盟管理", "多仓联盟、联盟订单、公众池协作。"),
            ("极速核销台", "扫码入库、扫码出库、盘点、顺丰码出入库匹配。"),
        ],
        widths=[2200, 7160],
    )

    doc.add_heading("6. 核心业务流程", level=1)
    doc.add_heading("6.1 订单生命周期", level=2)
    add_numbered(
        doc,
        [
            "创建订单：手工录入、一键复制解析、图片识别预留；按品牌型号或套餐录入客户、租期、渠道、租金、备注、需求、单号、视频等信息。",
            "前置校验：根据发货仓、租期范围、可接单库存、物流时效和天气风险判断是否可接单；库存不足时建议其他仓库或待选仓。",
            "备货发货：选择或扫码绑定 SN，自动匹配或人工上传发货视频，自动/手工下快递单，打印面单或订单码，确认发货。",
            "租用跟踪：订单进入租用中、发货在途、归还中、已签收、晚到、未知等状态，并按发货/归还/到货日期生成今日、明日、后日列表。",
            "归还验机：可直接确认收货，也可上传验机图片、记录划痕新旧、填写赔付金额、转售后或继续出租。",
            "完结与回补：确认收货或售后完成后，设备和配件按规则回补库存；外部平台是否完结可单独标记。",
            "取消/追回：未发货订单可取消且保持原排序；已发货订单只能追回，需填写原因并向仓库告警。",
        ],
    )
    add_table(
        doc,
        ["订单状态", "触发条件", "库存影响"],
        [
            ("待发货", "订单已创建，发货日未到或未确认发货。", "按租期占用可接单库存，SN 可未绑定。"),
            ("已备货", "SN 绑定完成但未确认发货。", "SN 状态变为已占用。"),
            ("发货在途", "已填写单号并确认发货。", "设备离开发货仓，占用状态持续。"),
            ("租用中", "客户已签收或租期中。", "不计入可用库存。"),
            ("归还中", "客户提交归还物流或系统抓取归还节点。", "仍不计入可用库存，纳入预计到库。"),
            ("逾期归还", "到期未归还，按平台默认或会员自定义规则触发。", "继续视为租用中，不释放库存。"),
            ("售后中", "验机发现损坏、缺失或赔付事项。", "根据继续出租、赔付、寄回、丢失规则更新。"),
            ("已完结", "确认收货或售后完成。", "设备状态回到在库可售或转出/维修/报损。"),
        ],
        widths=[1500, 4300, 3560],
        header_fill=LIGHT_BLUE,
    )

    doc.add_heading("6.2 库存与仓库流程", level=2)
    add_bullets(
        doc,
        [
            "库存分为在库空闲、发货在途、租用中、归还在途、设备报修、逾期归还、出库商品、续租商品等状态。",
            "预计库存按未来 30 天每日独立重新核算，不沿用前一日数值，避免误差累积。",
            "预计库存公式：总设备数 - 当日待发货数量 + 预计到库数量 - 租用中数量；逾期归还设备统一归入租用中。",
            "物流库存基于带归还单号设备与当前在库设备，按上午、下午、晚上预计到库时段拆分。",
            "每个库存数字应可穿透查看占用订单，方便核对账实差异。",
            "商品或配件占用超过总库存 90% 时标红预警；会员可按百分比或个数配置单品告警阈值。",
        ],
    )
    add_table(
        doc,
        ["场景", "系统动作"],
        [
            ("仓库库存不足", "自动提示有货仓库，给出 2-3 个调货/改仓建议。"),
            ("调货指令", "发起到目标仓，发货仓上传 SN 和快递单号，库存显示为发货在途。"),
            ("调货签收", "物流签收后 SN 自动迁移到达仓。"),
            ("归还仓改动", "客户 SN 自动切换至新归还仓并提醒新仓负责人；误收到货时可申请更新。"),
            ("联盟库存", "只共享 SN 库存，不共享套餐库存。"),
            ("盘点异常", "正常项静默留存；非在库但实物在场时弹窗登记工单/订单与备注。"),
        ],
        widths=[2200, 7160],
    )

    doc.add_heading("6.3 快递与客户归还流程", level=2)
    add_bullets(
        doc,
        [
            "会员下快递单需预充值，余额下限 100 元，单次支付不限制金额；可按客户配置 9 折或 8 折等折扣。",
            "系统需对接顺丰、京东、快递 100；前期重点支持顺丰和京东，控制接口查询频率。",
            "自动下单可设置在上门前 2 小时触发，人工可提前下单或取消，已下单后仍需支持修改上门时间。",
            "客户通过微信小程序归还时无需注册，授权或验证码确认手机号后匹配订单，系统锁定归还仓库并自动带出收寄件信息。",
            "寄件订单需在会员端、客户端和后台全链路展示待揽收、运输、签收、异常等状态，异常件标红预警。",
            "闲鱼归还单号与系统单号冲突时，以系统已有单号为准，不再覆盖。"
        ],
    )

    doc.add_heading("7. 主要功能模块设计", level=1)
    module_rows = [
        ("总后台", "会员使用情况、财务汇总、增值服务收入、快递费收入、平台支出、票据、总库存、城市订单量、客户脱敏、微信小程序授权与转发统计。"),
        ("首页", "今日任务、汇总报表、营收分析、租用中统计、租赁地图，按全部与各仓维度展示。"),
        ("订单管理", "录单、时效、SN、视频、物流、转发、转寄、验机、售后、逾期、取消追回、调货建议、工作台。"),
        ("仓库管理", "多维库存看板、仓库分类统计、可接单计算、库存保护、告警、明细穿透、未来库存预测。"),
        ("商品管理", "商品/配件/套餐、SN 档案、二维码条码、PT3 专属规则、套餐可租余量、成本收益与回报率。"),
        ("极速核销台", "统一扫码框，按页面区分入库、出库、盘点、顺丰码匹配，减少仓库操作步骤。"),
        ("公众号", "30 天免登录与活跃续期、逾期提醒、库存不足、时效错误、周转市场提醒、简易查看和拍 SN 完结。"),
        ("个人中心", "账务、工具、时效查询、在线表格、记账簿、二维码条码生成、记事本、仓库产品套餐与权限设置。"),
        ("周转市场", "库存展示、闲置出售、闲置出租、收购求购、站内私信、举报与保证金、排序与防作弊。"),
        ("多仓联盟", "联盟库存查看、协助订单、接单上限、接受/转派/拒绝发布公众池、四种发货方式。"),
        ("微信小程序", "客户订单查询、手机号验证、协议勾选、归还下单、订单卡片、归还海报、开箱视频上传。"),
        ("支付宝小程序", "多个小程序商户可绑定一个库存系统。"),
    ]
    add_table(doc, ["模块", "功能设计"], module_rows, widths=[1600, 7760], header_fill=LIGHT_BLUE)

    doc.add_heading("8. 商品、SN 与套餐模型", level=1)
    add_table(
        doc,
        ["对象", "关键字段/规则"],
        [
            ("商品", "品牌、型号、分类、状态、采购成本、出售售价、商品图片、详情描述、统计口径、是否前台展示、排序。"),
            ("SN 设备", "SN、设备照片、新旧程度、购买时间、激活时间、保修期、保险到期、当前仓库、当前状态、关联订单、发货视频使用频次。"),
            ("配件", "支持 SN 模式和简易总数模式；总数调大自动新增 ID，调小需手工删除并记录日志。"),
            ("套餐", "录单快捷模板，按主机与配件配比实时计算可租套数；可按平台创建，也可不选平台创建。"),
            ("PT3 主机", "强制逐条录入 SN，一级/二级分类系统默认，会员可新增三级分类。"),
            ("PT3 配件", "官方标配配件支持双录入，非官方配件归类为 PT3 渠道配件，通用配件可设置跨型号通用。"),
            ("二维码/条码", "每个商品自动生成永久二维码和条形码，可自定义下载尺寸，二维码下方带编号。"),
        ],
        widths=[1800, 7560],
    )
    add_callout(
        doc,
        "套餐计算原则",
        "套餐不是独立实物库存，而是基于主机 SN 与配件池配比计算出来的可租能力；下单选择套餐时应同步占用对应配件，剩余未占用配件作为备用库存展示。",
    )

    doc.add_heading("9. 数据对象与关系", level=1)
    add_table(
        doc,
        ["数据表/对象", "核心字段", "关系说明"],
        [
            ("Member", "会员类型、认证状态、昵称、主体、微信二维码、VIP 状态、空间额度、推广码", "拥有仓库、商品、订单、权限组、套餐、账务。"),
            ("Warehouse", "仓库名称、地址、负责人、上门规则、权限范围、是否联盟共享", "承接订单发货/归还，持有 SN 与配件库存。"),
            ("Product", "分类、品牌、型号、图片、描述、成本、统计配置", "挂接 SN、配件、套餐与市场帖子。"),
            ("AssetSN", "SN、状态、仓库、档案信息、照片、保修/保险、履历", "绑定订单明细、调货记录、盘点记录。"),
            ("AccessoryStock", "配件名称、模式、总数、可用数、占用数、日志", "被套餐配比和订单占用。"),
            ("PackageTemplate", "套餐名、平台、配件配比、适用商品、展示规则", "用于快速录单和套餐可租余量计算。"),
            ("RentalOrder", "客户、渠道、租期、仓库、SN、套餐、租金、状态、备注、编辑记录", "驱动库存占用、物流、视频、售后、账务。"),
            ("LogisticsOrder", "快递公司、单号、取件时间、路由节点、异常状态、费用、折扣", "绑定发货或归还订单。"),
            ("MediaAsset", "类型、文件、识别 SN、订单关联、保存策略、上传方", "承载发货视频、收货视频、开箱视频和验机图片。"),
            ("AfterSalesCase", "问题类型、赔付金额、缺失配件、处理进度、联系状态", "由验机订单转入，完成后回到可完结订单。"),
            ("AllianceOrder", "发起方、承接方、结算价、脱敏客户信息、发货方式、状态", "不在线支付，只记录协作和结算。"),
            ("MarketPost", "板块、品类、型号、成色、数量、城市、标签、价格区间、状态", "周转市场标准化展示，站内私信沟通。"),
            ("FinanceRecord", "收入/支出类型、金额、订单号、支付渠道、发票状态", "支持总后台收入、支出、票据与盈亏汇总。"),
        ],
        widths=[1700, 4200, 3460],
        header_fill=LIGHT_BLUE,
    )

    doc.add_heading("10. 外部接口与集成", level=1)
    add_table(
        doc,
        ["接口/渠道", "用途", "设计要求"],
        [
            ("顺丰接口", "下单、路由、签收、归还单匹配、扫码识别", "控制查询频率，全站 1 秒不超过 29 次；超频状态显示查询中，下秒继续。"),
            ("京东接口", "下单与路由", "与顺丰同类能力，作为前期重点。"),
            ("快递100", "路由查询与扩展", "预留接口，支持后期新增。"),
            ("微信小程序", "客户归还、订单查询、视频查看/上传", "不可搜索；通过订单卡片或二维码进入；授权和验证码均需说明用途。"),
            ("公众号", "提醒、简易管理、极速核销入口", "30 天不用重新登录，活跃自动续期。"),
            ("闲鱼/淘宝/抖音/美团/京东等渠道", "订单来源与外部订单状态", "渠道可由会员设置展示项；店铺名称变更不影响来源归集。"),
            ("支付宝租赁小程序", "后续租赁小程序接入", "支持多个小程序商户绑定同一个库存系统。"),
            ("免押码/免押服务", "订单免押能力", "可先生成免押码再补录订单，也可免押成功后同步客户订单和协议。"),
        ],
        widths=[1900, 3000, 4460],
    )

    doc.add_heading("11. 通知、风控与日志", level=1)
    add_table(
        doc,
        ["类别", "规则"],
        [
            ("业务提醒", "归还逾期、发货逾期、库存不足、时效错误、周转市场上新、保险到期前 1 个月提醒。"),
            ("订单日志", "所有订单需记录编辑历史、批注、备注、抽查记录、取消/追回原因。"),
            ("库存核查", "按 SN 流转、订单记录、近 3 个月使用记录做库存评分；总后台可按规则对会员库存评级。"),
            ("隐私脱敏", "联盟订单姓名隐藏 1 位、地址隐藏详细内容、手机号隐藏 6 位；微信小程序客户后台分权限查看。"),
            ("周转市场风控", "实物认证 SN 才可上架；限制重复上架；空挂库存降权；举报入口；假库存扣 500 元保证金。"),
            ("内容管控", "周转市场发帖采用全下拉选择，不开放任意文字输入；商家自定义转发话术需审核并留快照。"),
            ("接口风控", "物流接口限频，异常件标红，长时间不揽收/退回/滞留触发预警。"),
        ],
        widths=[2000, 7360],
        header_fill=LIGHT_BLUE,
    )

    doc.add_heading("12. 收费、增值服务与推广体系", level=1)
    add_table(
        doc,
        ["项目", "规则"],
        [
            ("免费基础能力", "电脑端和微信端订单管理、收发趋势、库存、基础发货时效、联盟仓库、发信息、1G 数据存储。"),
            ("批量视频/图片识别", "新会员免费 3 个月；月 1-300 单 30 元/月或 300 元/年；月 300 单以上 60 元/月或 600 元/年。"),
            ("快递路由", "月 1-500 单免费；501-1000 单 199 元/月或 1900 元/年；1000 单以上 299 元/月或 2900 元/年，并附赠时效检查。"),
            ("免押功能", "作为增值服务模块，支持订单免押码生成和信息同步。"),
            ("快递/物流差价", "按实际下单仓统计支出；物流赔付金额自动识别单号并退至会员。"),
            ("推广奖励", "有效用户标准：90 天内使用达 240 小时并完成 300 单正常订单；奖励 100 元/个。"),
            ("推广佣金", "被推广会员首次购买月/年费产品折扣后 50% 佣金，之后购买月/年费产品 20% 佣金。"),
            ("推广员资格", "需正规推广员邀约、注册满 15 天、累计在线不少于 80 小时、至少 50 笔真实成交、主体唯一、VIP 账户。"),
        ],
        widths=[2200, 7160],
    )

    doc.add_heading("13. 版本范围与实施优先级", level=1)
    add_table(
        doc,
        ["阶段", "建议范围", "验收重点"],
        [
            ("MVP/1.0", "首页、订单管理、仓库管理、商品/SN/套餐、收发趋势、个人中心基础设置、极速核销台、公众号提醒基础。", "能完成从录单、发货、归还、验机、完结到库存回补的闭环。"),
            ("1.1", "订单中转大厅、库存周转市场基础展示与三大板块标准化发布。", "坚持不定价、不交易、不担保；帖子标准化与私信规则可控。"),
            ("1.2", "多仓联盟、快递下单、微信小程序客户归还。", "联盟订单脱敏、快递全链路状态、客户无需注册即可归还。"),
            ("1.3", "免押码。", "支持先生成后补录，以及免押成功后同步订单与协议。"),
            ("2.0", "支付宝租赁小程序与多商户绑定。", "多个小程序商户可绑定一个库存系统，库存口径一致。"),
            ("远期", "库存同步接口、第三方租赁平台库存调用与溢出订单承接。", "授权、扣减、回补、外部渠道风控与结算边界明确。"),
        ],
        widths=[1200, 5200, 2960],
        header_fill=LIGHT_BLUE,
    )

    doc.add_heading("14. 关键验收清单", level=1)
    add_bullets(
        doc,
        [
            "录入订单时能够结合租期、仓库、库存、物流时效判断是否可接单，并能推荐调仓方案。",
            "SN 从入库、占用、发货、租用、归还、售后、完结、维修/出售均可追溯。",
            "套餐库存由主机与配件真实库存动态计算，不能与实物库存脱节。",
            "仓库人员只看到授权仓库，联盟成员只看到脱敏订单和共享 SN 库存。",
            "客户小程序归还流程无需注册，手机号授权/验证码说明清晰，归还仓库锁定。",
            "发货视频、客户收货视频、开箱视频、验机图片按保存策略与订单自动关联。",
            "周转市场不得出现租金、手机号、微信、详细地址、自由广告文案或平台担保交易闭环。",
            "财务后台能区分增值服务收入、快递相关收入、平台支出、发票状态和当期净收支。",
        ],
    )

    doc.add_heading("15. 待澄清事项", level=1)
    add_table(
        doc,
        ["问题", "建议处理"],
        [
            ("订单渠道自动抓取范围", "先明确闲鱼、淘宝等渠道是否具备可用接口或需人工/半自动导入。"),
            ("视频识别 SN 准确率", "MVP 可先做批量上传与人工复核，识别能力分阶段上线。"),
            ("时效历史匹配算法", "先采用接口匹配，沉淀历史数据后再切换或混合系统匹配。"),
            ("周转市场保证金扣罚流程", "需明确举报证据、人工审核、申诉、扣罚与限制入驻规则。"),
            ("快递折扣与会员奖励", "同一快递单是否客户 9 折和会员 10% 奖励互斥，需要财务策略确认。"),
            ("支付宝多商户绑定", "需确认一个库存系统下不同小程序商户的订单、财务、客户数据隔离规则。"),
        ],
        widths=[3000, 6360],
    )

    doc.add_page_break()
    doc.add_heading("附录 A：模块来源映射", level=1)
    add_table(
        doc,
        ["设计章节", "主要来源文档"],
        [
            ("总后台、财务、库存评分", "1.0版-0-总后台.docx"),
            ("导航与信息架构", "1.0版-1-栏目导航.docx"),
            ("订单生命周期", "ok-1.0版-3-订单管理.docx、1.1版-订单中转大厅.docx"),
            ("库存与仓库", "ok-1.0版-5-仓库管理.docx、1.0版-10-仓库调货.docx"),
            ("商品、SN、套餐", "ok-1.0版-6-商品管理-.docx"),
            ("极速核销", "ok-1.0版-9-极速核销台.docx"),
            ("微信/公众号/快递", "1.0版-7-公众号.docx、1.2版-快递下单及涉及流程+微信小程序.docx"),
            ("收费推广", "1.0版-15-免费与收费与推广.docx"),
            ("联盟与市场", "1.2版-多仓联盟.docx、1.1版-库存周转市场---待定.docx"),
            ("免押与支付宝小程序", "1.3版-免押码.docx、2.0版-支付宝租赁小程序.docx"),
        ],
        widths=[3000, 6360],
    )

    doc.save(DOCX_OUT)


def classify_source(name):
    if "总后台" in name:
        return "平台后台、财务、会员、统计"
    if "栏目导航" in name:
        return "信息架构"
    if "首页" in name:
        return "首页看板"
    if "订单管理" in name or "订单中转" in name:
        return "订单流程"
    if "收发趋势" in name:
        return "收发趋势"
    if "仓库管理" in name or "仓库调货" in name:
        return "仓库与库存"
    if "商品管理" in name:
        return "商品、SN、套餐"
    if "公众号" in name:
        return "公众号提醒与简易管理"
    if "个人中心" in name:
        return "会员设置与工具"
    if "核销" in name:
        return "扫码核销"
    if "收费" in name:
        return "商业化与推广"
    if "周转市场" in name:
        return "库存周转市场"
    if "多仓联盟" in name:
        return "多仓联盟"
    if "快递" in name or "微信小程序" in name:
        return "快递与客户归还小程序"
    if "免押" in name:
        return "免押能力"
    if "支付宝" in name:
        return "支付宝租赁小程序"
    return "其他"


def build_markdown():
    # The Markdown companion is intentionally concise and points to the DOCX as
    # the authoritative formatted deliverable.
    content = [
        "# 租赁项目设计文档",
        "",
        "本文件为同名 DOCX 的便于检索版本。正式排版交付件见：`租赁项目设计文档.docx`。",
        "",
        "## 核心定位",
        "",
        "迁渃租赁闭环及库存协同系统面向租赁商家，覆盖订单、库存、SN、仓库、物流、客户归还、售后完结、财务汇总、联盟协同和周转市场。",
        "",
        "## 核心模块",
        "",
    ]
    modules = [
        "总后台：会员、财务、库存、订单和小程序统计。",
        "订单管理：录单、时效、SN、视频、物流、转发、转寄、验机、售后、逾期、取消追回。",
        "仓库管理：多维库存、预计库存、物流库存、可接单数量、调货、盘点。",
        "商品管理：商品、配件、SN 档案、二维码条码、套餐与收益统计。",
        "极速核销台：扫码入库、扫码出库、盘点和快递码匹配。",
        "客户小程序：订单查询、视频查看、归还下单和开箱视频上传。",
        "多仓联盟：联盟共享 SN 库存、协助订单和公众池。",
        "库存周转市场：库存展示、标准化发帖、站内私信、举报与保证金。",
    ]
    for module in modules:
        content.append(f"- {module}")
    content.extend(["", "## 原始资料", ""])
    for i, name in enumerate(SOURCE_DOCS, 1):
        content.append(f"{i}. {name}")
    MD_OUT.write_text("\n".join(content), encoding="utf-8")


if __name__ == "__main__":
    build_docx()
    build_markdown()
