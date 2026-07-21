from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document


def iter_docx_text(path: Path) -> list[str]:
    doc = Document(path)
    lines: list[str] = []

    for paragraph in doc.paragraphs:
        text = " ".join(paragraph.text.split())
        if text:
            lines.append(text)

    for table_index, table in enumerate(doc.tables, start=1):
        lines.append(f"[表格 {table_index}]")
        for row in table.rows:
            cells = [" ".join(cell.text.split()) for cell in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))

    return lines


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source_dir", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()

    files = sorted(args.source_dir.glob("*.docx"))
    with args.output.open("w", encoding="utf-8") as out:
        out.write("# 租赁项目原始文档抽取汇总\n\n")
        for index, file in enumerate(files, start=1):
            out.write(f"## {index}. {file.name}\n\n")
            try:
                lines = iter_docx_text(file)
            except Exception as exc:  # noqa: BLE001
                out.write(f"> 抽取失败：{exc}\n\n")
                continue

            if not lines:
                out.write("> 未抽取到文本内容。\n\n")
                continue

            for line in lines:
                out.write(f"{line}\n\n")


if __name__ == "__main__":
    main()
