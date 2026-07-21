from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document

from build_rental_design_doc import (
    BLACK,
    LIGHT_BLUE,
    LIGHT_GRAY,
    MUTED,
    add_callout,
    add_para,
    add_table,
    style_doc,
)


OUT_DIR = Path("/Users/CharlesZhang/Documents/租赁系统")
DOCX_OUT = OUT_DIR / "租赁系统RuoYi-Cloud开发计划.docx"
MD_OUT = OUT_DIR / "租赁系统RuoYi-Cloud开发计划.md"


SOURCE_DOCS = [
    "租赁项目设计文档.docx",
    "租赁项目详细功能规格书.docx",
    "租赁项目原始文档抽取汇总.md",
]


ARCH_ROWS = [
    ("前端一：merchant-web", "商家端，基于 RuoYi-Cloud-Vue3-TypeScript 二次开发", "只允许商家管理本商户/主体下的订单、仓库、商品、库存、快递、客户归还、小程序配置、联盟协作等功能；所有接口必须自动带租户/商户上下文。"),
    ("前端二：admin-web", "平台内部管理后台，基于同框架独立工程或同 monorepo 独立 app", "仅平台内部人员使用；管理会员、商户认证、平台财务、风控、系统配置、周转市场审核、全局统计和运营工具。"),
    ("ruoyi-gateway", "统一 API 网关", "路由、鉴权前置、限流、黑名单、灰度、接口审计入口；按 merchant/admin 前端来源做策略区分。"),
    ("ruoyi-auth", "认证中心", "登录、刷新令牌、验证码、OAuth/微信授权预留；支持平台员工和商家员工两类身份。"),
    ("ruoyi-system", "系统基础服务", "用户、角色、菜单、部门、岗位、字典、参数、通知公告、操作日志；扩展商户、租户、仓库权限、敏感字段权限。"),
    ("ruoyi-file", "文件服务", "发货视频、验机图片、开箱视频、商品图片、设备图片；私有文件、临时访问 URL、大小/类型限制。"),
    ("ruoyi-job", "定时任务", "时效检查、逾期扫描、库存快照重算、保险到期提醒、佣金结算、视频过期清理、物流路由同步。"),
    ("ruoyi-gen", "代码生成", "用于基础 CRUD 快速生成；先生成字典、基础表、简单管理页；复杂流程和状态机由手写业务代码实现。"),
    ("rental-business", "租赁业务服务，建议新增微服务", "订单、商品、SN、套餐、仓库、库存、核销、售后、调货、收发趋势。"),
    ("rental-logistics", "物流服务，可先作为 rental-business 子模块", "快递下单、路由同步、时效检查、物流异常、客户归还小程序接口。"),
    ("rental-market", "库存周转市场服务，可后置", "市场帖子、筛选排序、站内私信、关注订阅、举报保证金、风控处罚。"),
    ("rental-finance", "财务/计费服务，可先模块化后独立", "会员余额、快递扣费、增值服务、佣金、保证金、发票和不可变流水。"),
]


REPO_ROWS = [
    ("后端仓库", "RuoYi-Cloud 原始结构 + 新增 rental-* 服务", "保留 ruoyi-auth/gateway/modules/common/api 等约定；新增业务服务不直接塞满 ruoyi-system。"),
    ("前端仓库方案 A", "两个独立前端工程 merchant-web、admin-web", "隔离最清晰，路由、菜单、权限、部署互不影响；推荐商业系统前期采用。"),
    ("前端仓库方案 B", "一个 monorepo 下两个 app", "共享组件和 API SDK 更方便，但构建配置稍复杂。若团队熟悉 pnpm workspace 可采用。"),
    ("共享前端包", "components/hooks/api-types/permission-utils", "抽出状态标签、表格筛选、上传组件、字典组件、权限指令、脱敏展示组件。"),
    ("接口契约", "OpenAPI/Knife4j + TypeScript 类型生成", "后端接口稳定后生成前端 API 类型，减少字段错配。"),
]


SECURITY_ROWS = [
    ("租户隔离", "所有业务表必须包含 tenant_id/member_id/merchant_id；MyBatis 查询默认拼接租户条件；平台后台跨租户查询必须显式授权并记录审计。"),
    ("双端权限", "商家端只读写本商户数据；平台端按内部角色划分运营、财务、风控、客服、系统管理员。前端隐藏菜单不能替代后端校验。"),
    ("仓库权限", "商家员工可被限制到指定仓库；订单、SN、视频、物流、盘点接口都要校验仓库范围。"),
    ("敏感信息", "客户姓名、手机号、地址、支付宝手机号、芝麻分、微信二维码默认脱敏；详情查看、导出、复制、转发均记录日志。"),
    ("资金流水", "余额、扣费、佣金、保证金、发票不可直接覆盖；通过流水、冲正、冻结、解冻实现。"),
    ("库存并发", "SN 绑定、库存预占、完结回补、调货迁移必须事务化；对同一 SN/订单做幂等和乐观锁/唯一约束。"),
    ("文件安全", "上传白名单、大小限制、MIME 校验、私有桶、临时 URL、防盗链；视频过期清理走定时任务。"),
    ("接口防护", "登录、验证码、物流查询、站内私信、上传、市场发布、举报、推广注册按用户/IP/设备/租户限流。"),
    ("审计日志", "权限变更、敏感查看、订单核心字段修改、SN 状态变化、库存扣减、财务变更、市场处罚必须有前后值。"),
]


PHASE_ROWS = [
    ("第 0 阶段：工程落地准备", "1 周", "确认 RuoYi-Cloud 版本、JDK/Node/MySQL/Redis/Nacos/Sentinel/Seata/MinIO 环境；建立后端服务命名、数据库命名、Git 分支、代码规范。", "本地能启动 gateway/auth/system；两个前端空壳可登录。"),
    ("第 1 阶段：双前端与权限骨架", "1-2 周", "搭建 merchant-web 与 admin-web；拆菜单、路由、布局、登录态、权限指令、字典、API 封装；后端扩展商户/租户/仓库权限。", "商家账号无法访问平台菜单；平台账号无法误进入商家工作台。"),
    ("第 2 阶段：基础资料与库存骨架", "2 周", "商品、品牌型号、仓库、配件、SN、套餐、二维码/条码、文件上传、库存流水表；先完成 CRUD 和权限。", "商家可维护本商户商品仓库；平台可查看汇总；跨商户数据不可见。"),
    ("第 3 阶段：订单主流程 MVP", "3-4 周", "添加订单、库存检测、SN 绑定、发货、单号、视频、归还验机、完结、取消/追回、订单列表和状态机。", "能完成录单到完结闭环；库存不负数；关键动作有审计。"),
    ("第 4 阶段：库存看板与仓库作业", "2-3 周", "在库空闲、预计库存、物流库存、库存穿透、告警、调货、极速核销台、盘点。", "仓库员工能完成扫码出入库；库存数字可追溯到订单/SN。"),
    ("第 5 阶段：物流与客户小程序接口", "2-3 周", "快递下单抽象、路由同步、时效检查、客户归还接口、小程序订单查询/开箱视频上传；前端先做 H5/接口联调页，后接微信小程序。", "客户只能看手机号匹配订单；归还仓锁定；物流异常可提醒。"),
    ("第 6 阶段：财务、增值服务、推广", "2-3 周", "会员余额、快递扣费、增值套餐、发票状态、佣金规则、推广员审核、VIP/存储额度。", "财务流水不可变；退款/冲正链路清楚。"),
    ("第 7 阶段：多仓联盟", "2 周", "联盟关系、共享 SN 库存、联盟协助订单、接单上限、转派/拒绝/公众池。", "联盟订单客户信息脱敏；只记账不支付。"),
    ("第 8 阶段：库存周转市场", "3 周", "标准化发帖、列表详情、筛选排序、私信、关注订阅、调拨意向、举报保证金、后台风控。", "不展示租金/私联；全下拉发布；违规可追溯。"),
    ("第 9 阶段：报表、运营与优化", "持续", "首页、收发趋势、复购、知识库、公告、平台总后台报表、性能优化、监控告警、安全加固。", "平台和商家均有可用看板；核心接口有监控。"),
]


MODULE_TASKS = [
    ("系统基础", "商户主体、认证状态、商户员工、角色、菜单、仓库权限、字典、参数、公告", "ruoyi-system 扩展 + admin/merchant 双端菜单"),
    ("商品中心", "品牌型号、商品、SN、配件、套餐、二维码条码、导入导出、成本收益字段", "rental-business: product 子模块"),
    ("仓库库存", "仓库、库存流水、库存快照、库存预占、调货、盘点、告警、库存穿透", "rental-business: inventory/warehouse 子模块"),
    ("订单中心", "录单、复制解析预留、订单状态机、SN 绑定、发货、归还、验机、售后、取消追回、工作台", "rental-business: order 子模块"),
    ("媒体中心", "发货视频、验机图片、客户开箱视频、商品图片、设备照片、识别 SN 预留", "ruoyi-file + MediaAsset 业务表"),
    ("物流中心", "承运商配置、下单、路由、时效、异常、顺丰码扫码匹配、快递余额扣费", "rental-logistics 或 business 子模块"),
    ("商家端首页", "今日待发/归还/到库、营收、租用中、租赁地图、工作提醒", "merchant-web dashboard"),
    ("平台总后台", "会员统计、财务汇总、总库存、订单统计、小程序授权、风控", "admin-web platform modules"),
    ("公众号/小程序接口", "30 天会话、提醒、客户订单查询、归还下单、开箱视频", "auth + logistics + order API"),
    ("多仓联盟", "联盟关系、共享库存、协助订单、转派、公众池", "rental-alliance 或 business 子模块"),
    ("周转市场", "帖子、标准字段、筛选、私信、举报、保证金、降权", "rental-market"),
    ("财务计费", "余额、扣费、增值服务、快递路由套餐、佣金、发票、保证金", "rental-finance"),
]


FRONTEND_ROWS = [
    ("登录与身份", "merchant-web 登录商家员工；admin-web 登录平台人员；两个前端使用不同路由白名单和菜单根。"),
    ("菜单策略", "商家端只放业务操作菜单；平台端放运营/风控/财务/系统配置；相同组件可共享但路由分离。"),
    ("API 封装", "按服务拆 api/order、api/product、api/inventory、api/logistics、api/finance；请求拦截器自动带 token 和端类型。"),
    ("权限指令", "复用 RuoYi v-hasPermi/v-hasRole，扩展仓库权限和敏感字段展示权限。"),
    ("商家端模块顺序", "首页、订单、发货/归还、仓库库存、商品/SN/套餐、极速核销、物流、客户小程序配置、联盟、市场、个人中心。"),
    ("平台端模块顺序", "总览、会员商户、认证审核、财务票据、总库存、订单统计、市场风控、推广佣金、系统管理、日志监控。"),
    ("共享组件", "状态标签、SN 选择器、仓库选择器、商品套餐选择器、视频上传、脱敏文本、审计日志抽屉、库存穿透弹窗。"),
]


BACKEND_ROWS = [
    ("包命名", "com.ruoyi.rental.order/product/inventory/logistics/finance/market/alliance", "保持 RuoYi Controller/Service/Mapper/Domain 分层。"),
    ("数据库", "ry_rental_order、ry_rental_product、ry_rental_asset_sn 等", "业务表统一 tenant_id/member_id + create_by/create_time/update_by/update_time。"),
    ("权限注解", "@PreAuthorize + 数据权限切面", "平台端和商家端接口分开权限标识，如 rental:order:add、platform:member:list。"),
    ("事务", "订单创建、库存预占、SN 绑定、完结回补、财务扣费", "必须 @Transactional；关键请求用幂等 key。"),
    ("消息/任务", "物流同步、视频清理、库存快照、逾期提醒", "先用 ruoyi-job，后续高并发再引 MQ。"),
    ("远程调用", "system/file/job 与 rental 服务", "使用 RuoYi Cloud OpenFeign 风格；外部接口统一封装 client。"),
    ("代码生成", "基础 CRUD 使用 ruoyi-gen", "复杂状态机、库存并发、财务流水不要完全依赖生成代码。"),
]


def title_page(doc: Document) -> None:
    add_para(doc, "RuoYi-Cloud 开发计划", size=24, bold=True, color=BLACK, after=4)
    add_para(doc, "迁渃租赁系统双前端 + 云服务二次开发任务拆分", size=14, color=MUTED, after=14)
    add_table(
        doc,
        ["项目", "说明"],
        [
            ("文档版本", "V1.0 开发计划版"),
            ("整理日期", date.today().isoformat()),
            ("技术基线", "RuoYi-Cloud-Vue3-TypeScript + RuoYi-Cloud 二次开发"),
            ("前端形态", "merchant-web 商家端 + admin-web 平台内部管理后台"),
            ("开发策略", "先搭整体模块框架、权限与数据隔离，再逐步完善模块详细功能。"),
        ],
        widths=[1800, 7560],
        header_fill=LIGHT_BLUE,
    )
    add_callout(
        doc,
        "核心拆分原则",
        "商家端是业务操作系统，只能管理本商户相关功能；平台后台是内部运营系统，用于跨商户监管、财务、风控和系统配置。后端必须以租户/商户隔离为第一约束，不能依赖前端菜单隐藏来保证安全。",
    )


def build_docx() -> None:
    doc = Document()
    style_doc(doc)
    title_page(doc)

    doc.add_heading("1. 现状与输入资料", level=1)
    add_para(doc, "当前工作区尚未放入 RuoYi-Cloud 源码，因此本计划按 RuoYi-Cloud 常见模块结构和已整理的租赁系统规格进行任务拆分。实际编码前需先导入或初始化 RuoYi-Cloud 与 RuoYi-Cloud-Vue3-TypeScript 工程。")
    add_table(doc, ["输入资料", "用途"], [(name, "功能范围与业务规则来源") for name in SOURCE_DOCS], widths=[3600, 5760], header_fill=LIGHT_GRAY)

    doc.add_heading("2. 总体架构拆分", level=1)
    add_table(doc, ["模块/工程", "定位", "二次开发内容"], ARCH_ROWS, widths=[1900, 2600, 4860], header_fill=LIGHT_BLUE)

    doc.add_heading("3. 仓库与工程组织建议", level=1)
    add_table(doc, ["事项", "建议", "说明"], REPO_ROWS, widths=[1900, 3200, 4260], header_fill=LIGHT_BLUE)

    doc.add_heading("4. 全局安全与数据隔离任务", level=1)
    add_table(doc, ["安全项", "开发要求"], SECURITY_ROWS, widths=[1900, 7460], header_fill=LIGHT_BLUE)

    doc.add_heading("5. 后端开发拆分", level=1)
    add_table(doc, ["任务", "建议做法", "注意事项"], BACKEND_ROWS, widths=[1900, 3900, 3560], header_fill=LIGHT_BLUE)

    doc.add_heading("6. 前端开发拆分", level=1)
    add_table(doc, ["任务", "开发要求"], FRONTEND_ROWS, widths=[2200, 7160], header_fill=LIGHT_BLUE)

    doc.add_heading("7. 业务模块任务清单", level=1)
    add_table(doc, ["业务模块", "首期任务", "建议落点"], MODULE_TASKS, widths=[1600, 5200, 2560], header_fill=LIGHT_BLUE)

    doc.add_heading("8. 分阶段开发计划", level=1)
    add_table(doc, ["阶段", "周期", "主要任务", "验收标准"], PHASE_ROWS, widths=[1350, 900, 5000, 2110], header_fill=LIGHT_BLUE)

    doc.add_heading("9. 前期模块骨架优先级", level=1)
    add_table(
        doc,
        ["优先级", "骨架内容", "原因"],
        [
            ("P0", "租户/商户、用户角色、菜单权限、仓库权限、字典、文件上传、审计日志", "所有业务模块都依赖这些基础能力，越晚做越容易返工。"),
            ("P0", "商品、仓库、SN、套餐、订单主表、库存流水的表结构和基础 CRUD", "先稳定领域模型，后续订单/库存/物流都围绕这些对象扩展。"),
            ("P1", "商家端与平台端的路由、菜单和权限分离", "双前端定位不同，必须一开始隔离，避免后期拆分成本高。"),
            ("P1", "订单状态机、SN 状态机、库存预占与回补接口", "这是租赁闭环的核心，不先做会影响所有页面真实联调。"),
            ("P2", "物流、视频、核销台、财务、联盟、市场", "这些模块可在主流程跑通后逐步接入。"),
        ],
        widths=[900, 5000, 3460],
        header_fill=LIGHT_BLUE,
    )

    doc.add_heading("10. 里程碑交付物", level=1)
    add_table(
        doc,
        ["里程碑", "交付物"],
        [
            ("M1 工程骨架", "后端服务可启动；merchant-web/admin-web 可登录；菜单权限和租户隔离生效。"),
            ("M2 基础资料", "商品、仓库、SN、套餐、文件上传、审计日志可用。"),
            ("M3 订单闭环", "订单创建、库存检测、SN 绑定、发货、归还、验机、完结可完整跑通。"),
            ("M4 仓库作业", "库存看板、调货、极速核销、盘点、告警可用。"),
            ("M5 外部协同", "物流、小程序接口、财务、联盟、周转市场逐步上线。"),
        ],
        widths=[1800, 7560],
        header_fill=LIGHT_BLUE,
    )

    doc.save(DOCX_OUT)


def build_markdown() -> None:
    lines = [
        "# 租赁系统 RuoYi-Cloud 开发计划",
        "",
        f"日期：{date.today().isoformat()}",
        "",
        "技术基线：RuoYi-Cloud-Vue3-TypeScript + RuoYi-Cloud 二次开发。",
        "",
        "前端拆分：merchant-web 商家端、admin-web 平台内部管理后台。",
        "",
    ]
    sections = [
        ("总体架构拆分", ARCH_ROWS),
        ("仓库与工程组织建议", REPO_ROWS),
        ("全局安全与数据隔离任务", SECURITY_ROWS),
        ("后端开发拆分", BACKEND_ROWS),
        ("前端开发拆分", FRONTEND_ROWS),
        ("业务模块任务清单", MODULE_TASKS),
        ("分阶段开发计划", PHASE_ROWS),
    ]
    for title, rows in sections:
        lines.append(f"## {title}")
        lines.append("")
        for row in rows:
            lines.append(f"- **{row[0]}**：{'；'.join(str(x) for x in row[1:])}")
        lines.append("")
    MD_OUT.write_text("\n".join(lines), encoding="utf-8")


if __name__ == "__main__":
    build_docx()
    build_markdown()
